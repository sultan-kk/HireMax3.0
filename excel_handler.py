"""
excel_handler.py
----------------
Writes verified candidate records + an audit trail of deleted /
irrelevant items into a single workbook (Candidates_Master.xlsx) with
two sheets:

    1. "Candidates"            -> one row per saved candidate
    2. "Audit_Irrelevant_Logs" -> one row per irrelevant/deleted item

Also supports:
    - Deleting a saved candidate record (+ its audit rows) permanently.
    - Custom column display-names (via config.load_column_labels()).
    - Producing a nicely-formatted, ready-to-download copy of the
      workbook in-memory (borders, autofilter, frozen header, autofit
      column widths, phone column forced to TEXT so leading '+'/'0'
      isn't stripped by Excel).
"""

import io
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

import config

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="B7C6D9"),) * 4)


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------
def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22


def _current_headers():
    labels = config.load_column_labels()
    candidate_headers = [labels["candidate"][k] for k in config.CANDIDATE_FIELD_KEYS]
    audit_headers = [labels["audit"][k] for k in config.AUDIT_FIELD_KEYS]
    return candidate_headers, audit_headers


def _ensure_workbook(path: str):
    """Creates the workbook with both sheets + current headers if it doesn't exist yet."""
    if os.path.exists(path):
        return
    candidate_headers, audit_headers = _current_headers()

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Candidates"
    ws1.append(candidate_headers)
    _style_header(ws1)

    ws2 = wb.create_sheet("Audit_Irrelevant_Logs")
    ws2.append(audit_headers)
    _style_header(ws2)

    wb.save(path)


def get_next_candidate_id(path: str = config.MASTER_EXCEL_PATH) -> int:
    if not os.path.exists(path):
        return 1
    try:
        df = pd.read_excel(path, sheet_name="Candidates")
        id_col = config.load_column_labels()["candidate"]["candidate_id"]
        if df.empty or id_col not in df.columns:
            return 1
        return int(df[id_col].max()) + 1
    except Exception:
        return 1


def find_duplicate(email: str, phone: str, path: str = config.MASTER_EXCEL_PATH):
    """
    HR feature: checks whether a candidate with this email OR phone is
    already saved, so the reviewer gets a heads-up before creating an
    accidental duplicate record (e.g. the same CV uploaded twice, or two
    different files for the same applicant).
    Returns the existing Candidate ID if found, else None.
    """
    if not os.path.exists(path):
        return None
    try:
        df = pd.read_excel(path, sheet_name="Candidates")
    except Exception:
        return None
    if df.empty:
        return None

    labels = config.load_column_labels()["candidate"]
    id_col = labels["candidate_id"]
    email_col = labels["email"]
    phone_col = labels["phone"]

    email = (email or "").strip().lower()
    phone_digits = "".join(ch for ch in (phone or "") if ch.isdigit())

    for _, row in df.iterrows():
        row_email = str(row.get(email_col, "")).strip().lower()
        row_phone_digits = "".join(ch for ch in str(row.get(phone_col, "")) if ch.isdigit())
        if email and row_email == email:
            return row[id_col]
        # Compare last 9 digits so "+92 300 1234567" matches "0300-1234567"
        if phone_digits and row_phone_digits and phone_digits[-9:] == row_phone_digits[-9:]:
            return row[id_col]
    return None


# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
def save_candidate(
    candidate_data: dict,
    irrelevant_items: list,
    source_file_name: str,
    saved_by: str = "HR Reviewer",
    path: str = config.MASTER_EXCEL_PATH,
) -> int:
    """
    Appends one candidate row + N audit rows to the master workbook.
    Returns the Candidate ID assigned to this row.
    """
    _ensure_workbook(path)
    wb = load_workbook(path)
    ws_candidates = wb["Candidates"]
    ws_audit = wb["Audit_Irrelevant_Logs"]

    candidate_id = get_next_candidate_id(path)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    ws_candidates.append([
        candidate_id,
        candidate_data.get("name", ""),
        candidate_data.get("email", ""),
        str(candidate_data.get("phone", "")),  # keep as text — preserves '+', leading 0
        candidate_data.get("skills", ""),
        candidate_data.get("experience_years", ""),
        candidate_data.get("education", ""),
        source_file_name,
        timestamp,
        saved_by,
    ])

    for item in irrelevant_items:
        ws_audit.append([
            candidate_id,
            candidate_data.get("name", ""),
            item.get("text", ""),
            item.get("reason", ""),
            item.get("confidence", ""),
            item.get("status", "deleted"),
            timestamp,
        ])

    _style_header(ws_candidates)
    _style_header(ws_audit)
    wb.save(path)
    return candidate_id


# ---------------------------------------------------------------------------
# Delete a saved record (candidate row + its audit rows)
# ---------------------------------------------------------------------------
def delete_candidate(candidate_id, path: str = config.MASTER_EXCEL_PATH) -> bool:
    """
    Permanently removes the candidate row with this ID from "Candidates",
    and every audit row referencing that ID from "Audit_Irrelevant_Logs".
    Returns True if a candidate row was found and removed.
    """
    if not os.path.exists(path):
        return False

    wb = load_workbook(path)
    ws_candidates = wb["Candidates"]
    ws_audit = wb["Audit_Irrelevant_Logs"]

    id_col_idx = 1  # "Candidate ID" is always the first column
    deleted = False

    for row in list(ws_candidates.iter_rows(min_row=2))[::-1]:
        if row[id_col_idx - 1].value == candidate_id:
            ws_candidates.delete_rows(row[0].row, 1)
            deleted = True

    for row in list(ws_audit.iter_rows(min_row=2))[::-1]:
        if row[id_col_idx - 1].value == candidate_id:
            ws_audit.delete_rows(row[0].row, 1)

    if deleted:
        _style_header(ws_candidates)
        _style_header(ws_audit)
        wb.save(path)
    return deleted


def rewrite_column_headers(path: str = config.MASTER_EXCEL_PATH):
    """
    Overwrites row 1 of both sheets with the CURRENT custom column labels
    (called right after the user saves new column names in the UI), without
    touching any existing data rows.
    """
    if not os.path.exists(path):
        return
    candidate_headers, audit_headers = _current_headers()
    wb = load_workbook(path)
    for idx, header in enumerate(candidate_headers, start=1):
        wb["Candidates"].cell(row=1, column=idx, value=header)
    for idx, header in enumerate(audit_headers, start=1):
        wb["Audit_Irrelevant_Logs"].cell(row=1, column=idx, value=header)
    _style_header(wb["Candidates"])
    _style_header(wb["Audit_Irrelevant_Logs"])
    wb.save(path)


# ---------------------------------------------------------------------------
# Formatted download copy (borders, autofilter, freeze panes, autofit,
# phone-as-text) — built in-memory so the master file on disk is untouched.
# ---------------------------------------------------------------------------
def build_formatted_download(path: str = config.MASTER_EXCEL_PATH) -> bytes:
    """
    Loads the master workbook, applies presentation formatting (autofilter,
    frozen header row, cell borders, autofit-ish column widths, phone
    column forced to text format), and returns the result as raw .xlsx
    bytes suitable for st.download_button. The file on disk is NOT modified.
    """
    _ensure_workbook(path)
    wb = load_workbook(path)

    for ws in wb.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue

        _style_header(ws)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Borders + autofit width based on the longest value in each column
        col_widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.border = THIN_BORDER
                length = len(str(cell.value)) if cell.value is not None else 0
                col_widths[cell.column_letter] = max(col_widths.get(cell.column_letter, 10), length + 2)
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = min(width, 45)

        # Force the "Phone" column to TEXT format so Excel doesn't strip a
        # leading '+' / '0' or convert it to a number.
        labels = config.load_column_labels()
        phone_label = labels["candidate"].get("phone", "Phone")
        header_values = [c.value for c in ws[1]]
        if phone_label in header_values:
            phone_col_idx = header_values.index(phone_label) + 1
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=phone_col_idx).number_format = "@"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Read helpers (used by the "view summary" panel)
# ---------------------------------------------------------------------------
def load_all_candidates(path: str = config.MASTER_EXCEL_PATH) -> pd.DataFrame:
    if not os.path.exists(path):
        return pd.DataFrame(columns=_current_headers()[0])
    return pd.read_excel(path, sheet_name="Candidates")


def load_audit_log(path: str = config.MASTER_EXCEL_PATH) -> pd.DataFrame:
    if not os.path.exists(path):
        return pd.DataFrame(columns=_current_headers()[1])
    return pd.read_excel(path, sheet_name="Audit_Irrelevant_Logs")
